"""Document ingestion: store original → parse → chunk → embed → index → register.

Supports PDF, DOCX, and plain-text submissions, SOPs, templates, and
historical NOTAs.
"""
import io
import uuid

from app.domain.models import (
    DEFAULT_ENTITY,
    BPIEntity,
    Chunk,
    DataClassification,
    DocumentType,
    GovernanceDocument,
    IngestionResult,
    new_trace_id,
)
from app.domain.ports import AuditLog, CaseRepository, Embedder, ObjectStorage, VectorStore

CHUNK_SIZE = 1200
CHUNK_OVERLAP = 150


def parse_bytes(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if name.endswith(".docx"):
        import docx

        document = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in document.paragraphs)
    return data.decode("utf-8", errors="replace")


def chunk_text(content: str) -> list[str]:
    content = content.strip()
    if not content:
        return []
    chunks, start = [], 0
    while start < len(content):
        end = start + CHUNK_SIZE
        chunks.append(content[start:end])
        start = end - CHUNK_OVERLAP
    return chunks


async def ingest_document(
    filename: str,
    data: bytes,
    doc_type: DocumentType,
    case_id: str | None,
    *,
    embedder: Embedder,
    vectors: VectorStore,
    storage: ObjectStorage,
    repository: CaseRepository,
    audit: AuditLog,
    entity: BPIEntity = DEFAULT_ENTITY,
    classification: DataClassification = DataClassification.INTERNAL,
) -> IngestionResult:
    document_id = f"doc_{uuid.uuid4().hex[:12]}"
    trace_id = new_trace_id()

    # 1. Preserve the original file (evidence trail), namespaced by entity.
    storage_key = f"{entity.value}/cases/{case_id or 'unassigned'}/{document_id}/{filename}"
    storage.put(storage_key, data)

    # 2. Parse and index for retrieval (chunks carry the tenant scope).
    content = parse_bytes(filename, data)
    pieces = chunk_text(content)
    chunks = [
        Chunk(
            document_id=document_id,
            case_id=case_id,
            entity=entity,
            doc_type=doc_type.value,
            chunk_index=i,
            content=piece,
        )
        for i, piece in enumerate(pieces)
    ]
    embeddings = await embedder.embed_texts([c.content for c in chunks])
    indexed = await vectors.replace_document(chunks, embeddings)

    # 3. Register document metadata (submissions are the master record).
    await repository.save_document(
        GovernanceDocument(
            id=document_id,
            case_id=case_id,
            entity=entity,
            classification=classification,
            doc_type=doc_type,
            title=filename,
            is_master=doc_type == DocumentType.SUBMISSION,
            storage_key=storage_key,
        )
    )

    audit.log(
        "document.ingested",
        trace_id,
        {
            "document_id": document_id,
            "filename": filename,
            "doc_type": doc_type.value,
            "case_id": case_id,
            "entity": entity.value,
            "classification": classification.value,
            "storage_key": storage_key,
            "chunks_indexed": indexed,
        },
    )
    return IngestionResult(document_id=document_id, chunks_indexed=indexed,
                           storage_key=storage_key)
